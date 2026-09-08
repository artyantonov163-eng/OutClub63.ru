#!/usr/bin/env python3
"""Reproduce the reviewed 2026-09-06 club snapshot; no network or chat export."""
import argparse
import io
import json
from pathlib import Path
import shutil

import openpyxl
from PIL import Image

# Schedule rows; May event spans rows 12..14. No dates inferred for planned events.
# Transcribed from the final screenshot 2565, visually checked. 2551 is an earlier
# partial version, never counted again. Only exact unique names are joined.
# Published historical order in screenshot 954, not a record of played scores.


def build(source, output):
    settings = json.loads((output / 'private/data/snapshot-settings.json').read_text())
    NAMES = settings['NAMES']
    TOTALS = settings['TOTALS']
    TITLES = settings['TITLES']
    SCHEDULE_ROWS = settings['SCHEDULE_ROWS']
    DOCS = {int(k): v for k, v in settings['DOCS'].items()}
    PHOTOS = {int(k): v for k, v in settings['PHOTOS'].items()}
    STRUCTURES = {int(k): v for k, v in settings['STRUCTURES'].items()}
    MATCHES = settings['MATCHES']
    DOUBLES_PAIRS = settings['DOUBLES_PAIRS']
    DOUBLES_SCHEDULE = settings['DOUBLES_SCHEDULE']
    workbook = openpyxl.load_workbook(source / 'table.xlsx', data_only=True)
    rating, schedule = workbook['Рейтинг'], workbook['Расписание']
    media = source / 'context-media'
    inventory = {x['message_id']: x for x in json.loads((media / 'inventory.json').read_text())}
    dest = output / 'private'
    for folder in ['data', 'media/players', 'media/documents', 'media/tournaments']:
        (dest / folder).mkdir(parents=True, exist_ok=True)
    images = {img.anchor._from.row + 1: img for img in rating._images}
    players = []
    for index, name in enumerate(NAMES):
        row = index + 5
        assert str(rating.cell(row, 4).value).strip() == name, 'Snapshot player order changed; review ID mapping first'
        assert int(rating.cell(row, 1).value) == index + 1
        pid = f'player-{index+1:02}'
        results = [{'tournamentId': f't{col-4:02}', 'points': int(rating.cell(row, col).value)} for col in range(5, 22) if isinstance(rating.cell(row, col).value, (float, int))]
        points = int(rating.cell(row, 2).value)
        assert points == sum(r['points'] for r in results) == TOTALS[index], f'Total mismatch: {pid}'
        photo = None
        if row in images:
            # Extract source photograph; conversion only, no generative alteration.
            image = Image.open(io.BytesIO(images[row]._data())).convert('RGB')
            image.save(dest / 'media' / 'players' / f'{pid}.jpg', quality=95)
            photo = f'/media/players/{pid}.jpg'
        players.append(dict(id=pid, name=name, rank=index+1, points=points, photo=photo, results=results))

    def copy_media(message_id, folder, filename):
        original = media / inventory[message_id]['local_file']
        assert original.is_file(), f'Missing reviewed attachment {message_id}'
        shutil.copyfile(original, dest / 'media' / folder / filename)
        return f'/media/{folder}/{filename}'

    tournaments = []
    for i, name in enumerate(TITLES):
        tid = f't{i+1:02}'
        completed = i < 12
        row = SCHEDULE_ROWS[i] if completed else None
        results = [{'playerId': p['id'], 'points': r['points']} for p in players for r in p['results'] if r['tournamentId'] == tid]
        event = dict(id=tid, name=name, date=schedule.cell(row, 2).value.date().isoformat() if completed else None,
            category=('OCT ' + ('500' if '500' in schedule.cell(row, 6).value else '250')) if completed else 'Категория уточняется',
            format=('doubles' if i in [2, 5, 8, 11] else 'singles') if completed else None, status='completed' if completed else 'planned', rated=True,
            venue=schedule.cell(row, 5).value if completed else 'Место уточняется',
            description='Результаты и очки из клубной таблицы. Организационные вопросы — в чате OUT.' if completed else 'Турнир указан в календаре сезона. Дата и подробности появятся после объявления в чате.',
            phases=[], participants=[r['playerId'] for r in results], results=results, documents=[], photos=[], matches=[],
            coverageNote='Показаны участники с начислениями в рейтинговой таблице; гости могут не входить в этот список. Полные счета матчей пока не восстановлены.' if completed else 'Дата ещё не опубликована. Запись и все договорённости остаются в чате.')
        if i in DOCS:
            message_id, filename = DOCS[i]
            event['documents'].append({'name': 'Регламент турнира', 'url': copy_media(message_id, 'documents', filename)})
        for message_id in PHOTOS.get(i, []):
            event['photos'].append({'url': copy_media(message_id, 'tournaments', f'event-{message_id}.jpg'), 'caption': f'{name} · фотографии клуба'})
        if i in STRUCTURES:
            event['bracket'] = [{'name': title, 'description': description} for title, description in STRUCTURES[i]]
        tournaments.append(event)
    doubles = tournaments[2]
    assert inventory[954]['date'].startswith('2026-02-13'), 'Unexpected date of doubles schedule'
    doubles['pairs'] = [{'names': names, 'sourceId': 954} for names in DOUBLES_PAIRS]
    doubles['schedule'] = [{'court': 'Корт 1' if index < 6 else 'Корты 1 и 2', 'time': time, 'sideA': DOUBLES_PAIRS[a], 'sideB': DOUBLES_PAIRS[b], 'sourceId': 954} for index, (time, a, b) in enumerate(DOUBLES_SCHEDULE)]
    doubles['coverageNote'] = 'Пары и порядок 10 игр сохранены из расписания, опубликованного перед турниром. Это план, не подтверждение сыгранных матчей. После 22:00 игры параллельные; точное распределение по кортам не указано. Прозвища сохранены как в источнике.'
    summer = tournaments[7]
    summer['endDate'] = '2026-05-16'
    summer['phases'] = [{'date': '2026-05-09', 'label': 'Групповой этап'}, {'date': '2026-05-10', 'label': 'Полуфиналы'}, {'date': '2026-05-16', 'label': 'Финалы'}]
    exact_names = {p['name']: p['id'] for p in players if sum(other['name'] == p['name'] for other in players) == 1}
    for index, (side_a, side_b, score) in enumerate(MATCHES):
        match = dict(id=f't08-m{index+1:02}', sideA=[side_a], sideB=[side_b], score=score, stage=f'Групповой этап · группа {index//3+1}', sourceId=2565)
        if side_a in exact_names:
            match['playerAIds'] = [exact_names[side_a]]
        if side_b in exact_names:
            match['playerBIds'] = [exact_names[side_b]]
        summer['matches'].append(match)
    # Source labels only: disjoint complete three-player round robins identify
    # four groups without asserting any alias-to-profile identity or draw name.
    summer['groups'] = []
    for group_index in range(4):
        source_matches = MATCHES[group_index*3:group_index*3+3]
        names = list(dict.fromkeys(name for a, b, _ in source_matches for name in [a, b]))
        assert len(names) == 3
        standings = {name: dict(name=name, played=0, wins=0, losses=0, gamesFor=0, gamesAgainst=0) for name in names}
        for a, b, score in source_matches:
            games_a, games_b = map(int, score.split(':'))
            assert games_a != games_b
            for name, own, other in [(a, games_a, games_b), (b, games_b, games_a)]:
                stats = standings[name]
                stats['played'] += 1
                stats['wins'] += int(own > other)
                stats['losses'] += int(own < other)
                stats['gamesFor'] += own
                stats['gamesAgainst'] += other
        assert all(s['played'] == 2 for s in standings.values())
        ordered = sorted(standings.values(), key=lambda s: -s['wins'])
        assert [s['wins'] for s in ordered] == [2, 1, 0]
        summer['groups'].append({'name': f'Группа {group_index+1}', 'players': names, 'standings': ordered})
    summer['bracket'][1]['participants'] = settings['GOLD_PARTICIPANTS']
    summer['bracket'][2]['participants'] = settings['SILVER_PARTICIPANTS']
    summer['coverageNote'] = 'Группы восстановлены по записи 12 матчей; названия условные. Прозвища сохранены как в источнике и не все сопоставлены с профилями. Участники золотой и серебряной лиг указаны на листе результатов; пары полуфиналов и счета плей-офф пока не восстановлены.'
    data = dict(schemaVersion=1, updatedAt='2026-09-06T00:00:00+04:00', season=2026, title='OUT Club Tour', players=players, tournaments=tournaments,
        rules={'summary': ['Общий зачёт включает одиночные и парные турниры.', 'Показаны опубликованные очки клубной таблицы.', 'Формат матчей и сеток определяется регламентом конкретного турнира.', 'Запись, отмены, жеребьёвка и договорённости остаются в чате OUT.'],
            'ratingNote': 'В таблице суммируются все начисления. Сезонный PDF описывает девять лучших результатов; до уточнения сохраняем опубликованные итоги таблицы.'},
        coverage={'players': len(players), 'photos': sum(p['photo'] is not None for p in players), 'tournaments': len(tournaments), 'completedTournaments': sum(t['status'] == 'completed' for t in tournaments), 'matches': len(MATCHES),
            'notes': ['Данные на 6 сентября 2026 года. Это сохранённый снимок, без автоматического обновления.', 'В рейтинге 16 игроков; полный состав клуба может быть больше.', 'Одноимённые участники сохраняют отдельные идентификаторы.', 'Счета найдены для 12 групповых матчей одного турнира. Это не полная статистика сезона.', 'Будущие даты и составы уточняются в чате.']})
    import runpy
    enrich = runpy.run_path(str(Path(__file__).with_name('enrich-presentation.py')))['enrich']
    data = enrich(data, output)
    (dest / 'data' / 'club.json').write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(f'Prepared {len(players)} players, {len(tournaments)} tournaments, {len(MATCHES)} matches; all published totals verified.')


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root', type=Path, default=root.parent / 'tennis-club-out')
    parser.add_argument('--output-root', type=Path, default=root)
    args = parser.parse_args()
    build(args.source_root.resolve(), args.output_root.resolve())
